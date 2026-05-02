"""
研发部门 API：调色板色块 CRUD（每个用户自己的常用色，存 PostgreSQL）
包含：调色板 / 历史色号 / AI 配色生成
"""
import uuid
from typing import Optional

from fastapi import APIRouter, HTTPException, Header, UploadFile, File
from pydantic import BaseModel, Field

from backend.db import SessionLocal, is_db_configured
from backend.db.models import RdColorSwatch, RdSwatch
from backend.auth import decode_token

router = APIRouter(prefix="/rd", tags=["rd"])


def _get_user_id(authorization: Optional[str]) -> str:
    """从 Authorization Bearer token 提取 user_id（username），未登录时抛 401"""
    if not authorization or not authorization.startswith("Bearer "):
        raise HTTPException(status_code=401, detail="未登录")
    token = authorization.split(" ", 1)[1]
    payload = decode_token(token)
    if not payload:
        raise HTTPException(status_code=401, detail="token 无效或已过期")
    user_id = payload.get("sub")
    if not user_id:
        raise HTTPException(status_code=401, detail="token 缺少用户信息")
    return user_id


# ---------- Schemas ----------

class ColorSwatchCreate(BaseModel):
    hex: str = Field(..., description="颜色值，如 #FF6B9D")
    label: str = Field("", description="可选备注名")


class ColorSwatchOut(BaseModel):
    id: str
    hex: str
    label: str
    sort_order: int
    created_at: str

    class Config:
        from_attributes = True


# ---------- Routes ----------

@router.get("/color-swatches")
async def list_swatches(authorization: Optional[str] = Header(None)):
    """获取当前用户的调色板色块列表"""
    if not is_db_configured():
        return []
    user_id = _get_user_id(authorization)
    db = SessionLocal()
    try:
        rows = (
            db.query(RdColorSwatch)
            .filter(RdColorSwatch.user_id == user_id)
            .order_by(RdColorSwatch.sort_order.asc(), RdColorSwatch.created_at.asc())
            .all()
        )
        return [
            {
                "id": r.id,
                "hex": r.hex,
                "label": r.label or "",
                "sort_order": r.sort_order,
                "created_at": r.created_at.isoformat() if r.created_at else "",
            }
            for r in rows
        ]
    finally:
        db.close()


@router.post("/color-swatches", status_code=201)
async def create_swatch(body: ColorSwatchCreate, authorization: Optional[str] = Header(None)):
    """添加一个色块到当前用户的调色板"""
    if not is_db_configured():
        raise HTTPException(status_code=503, detail="数据库未配置")
    user_id = _get_user_id(authorization)

    # 校验 HEX 格式（宽松：3/6/8 位）
    hex_val = body.hex.strip()
    if not hex_val.startswith("#") or len(hex_val) not in (4, 7, 9):
        raise HTTPException(status_code=422, detail="颜色值格式不正确，需要 #RGB / #RRGGBB / #RRGGBBAA")

    db = SessionLocal()
    try:
        # 最多存 32 个
        count = db.query(RdColorSwatch).filter(RdColorSwatch.user_id == user_id).count()
        if count >= 32:
            raise HTTPException(status_code=400, detail="调色板最多保存 32 个色块")

        row = RdColorSwatch(
            id=str(uuid.uuid4()),
            user_id=user_id,
            hex=hex_val.upper(),
            label=body.label or "",
            sort_order=count,
        )
        db.add(row)
        db.commit()
        db.refresh(row)
        return {
            "id": row.id,
            "hex": row.hex,
            "label": row.label,
            "sort_order": row.sort_order,
            "created_at": row.created_at.isoformat() if row.created_at else "",
        }
    finally:
        db.close()


@router.delete("/color-swatches/{swatch_id}", status_code=204)
async def delete_swatch(swatch_id: str, authorization: Optional[str] = Header(None)):
    """删除调色板中的一个色块（只能删自己的）"""
    if not is_db_configured():
        raise HTTPException(status_code=503, detail="数据库未配置")
    user_id = _get_user_id(authorization)
    db = SessionLocal()
    try:
        row = db.query(RdColorSwatch).filter(
            RdColorSwatch.id == swatch_id,
            RdColorSwatch.user_id == user_id
        ).first()
        if not row:
            raise HTTPException(status_code=404, detail="色块不存在或无权限删除")
        db.delete(row)
        db.commit()
    finally:
        db.close()


# ============================================================
# 研发历史色号记录（左侧历史列表）
# ============================================================

TEXTURES = ['哑光', '缎面', '镜面', '金属']


class RdSwatchCreate(BaseModel):
    name: str = Field(..., description="色号名称")
    hex: str = Field(..., description="颜色值")
    texture: str = Field("哑光", description="质地")
    opacity: int = Field(100, ge=0, le=100, description="透明度 0-100")
    note: str = Field("", description="备注")
    session_id: str = Field("", description="关联的趋势报告 session_id")


@router.get("/swatches")
async def list_swatches_history(authorization: Optional[str] = Header(None)):
    """获取当前用户的研发历史色号列表（左侧历史栏）"""
    if not is_db_configured():
        return []
    user_id = _get_user_id(authorization)
    db = SessionLocal()
    try:
        rows = (
            db.query(RdSwatch)
            .filter(RdSwatch.user_id == user_id)
            .order_by(RdSwatch.created_at.desc())
            .all()
        )
        return [
            {
                "id": r.id,
                "name": r.name,
                "hex": r.hex,
                "texture": r.texture or "哑光",
                "opacity": r.opacity,
                "note": r.note or "",
                "session_id": r.session_id or "",
                "date": r.created_at.strftime("%Y-%m-%d") if r.created_at else "",
            }
            for r in rows
        ]
    finally:
        db.close()


@router.post("/swatches", status_code=201)
async def create_swatch_history(body: RdSwatchCreate, authorization: Optional[str] = Header(None)):
    """保存一条色号记录到研发历史（点击「保存配方」时）"""
    if not is_db_configured():
        raise HTTPException(status_code=503, detail="数据库未配置")
    user_id = _get_user_id(authorization)

    hex_val = body.hex.strip()
    if not hex_val.startswith("#") or len(hex_val) not in (4, 7, 9):
        raise HTTPException(status_code=422, detail="颜色值格式不正确")
    if body.texture not in TEXTURES:
        raise HTTPException(status_code=422, detail=f"质地必须是 {TEXTURES} 之一")

    db = SessionLocal()
    try:
        row = RdSwatch(
            id=str(uuid.uuid4()),
            user_id=user_id,
            name=body.name.strip() or "色号",
            hex=hex_val.upper(),
            texture=body.texture,
            opacity=body.opacity,
            note=body.note or "",
            session_id=body.session_id or "",
        )
        db.add(row)
        db.commit()
        db.refresh(row)
        return {
            "id": row.id,
            "name": row.name,
            "hex": row.hex,
            "texture": row.texture,
            "opacity": row.opacity,
            "note": row.note,
            "session_id": row.session_id,
            "date": row.created_at.strftime("%Y-%m-%d") if row.created_at else "",
        }
    finally:
        db.close()


@router.delete("/swatches/{swatch_id}", status_code=204)
async def delete_swatch_history(swatch_id: str, authorization: Optional[str] = Header(None)):
    """删除一条研发历史色号（只能删自己的）"""
    if not is_db_configured():
        raise HTTPException(status_code=503, detail="数据库未配置")
    user_id = _get_user_id(authorization)
    db = SessionLocal()
    try:
        row = db.query(RdSwatch).filter(
            RdSwatch.id == swatch_id,
            RdSwatch.user_id == user_id
        ).first()
        if not row:
            raise HTTPException(status_code=404, detail="色号不存在或无权限删除")
        db.delete(row)
        db.commit()
    finally:
        db.close()


# ============================================================
# AI 配色生成（调用 agent generate_color_schemes 工具）
# ============================================================

class ColorSchemesRequest(BaseModel):
    session_id: str = Field(..., description="agent 会话 ID，用于获取已分析的色调数据")


@router.post("/color-schemes")
async def generate_color_schemes(body: ColorSchemesRequest, authorization: Optional[str] = Header(None)):
    """根据 agent 已分析的色调数据，调用 LLM 生成 3 组配色方案"""
    _get_user_id(authorization)  # 确认登录

    # 从已持久化的 agent 会话中获取 tone_analysis
    from backend.db import SessionLocal as _SL, is_db_configured as _idc
    from backend.db.models import AgentSession
    import json as _json

    if not _idc():
        raise HTTPException(status_code=503, detail="数据库未配置")

    db = _SL()
    try:
        sess = db.query(AgentSession).filter(AgentSession.id == body.session_id).first()
        if not sess:
            raise HTTPException(status_code=404, detail="找不到该 session")
        final_report = sess.final_report
        if isinstance(final_report, str):
            final_report = _json.loads(final_report)
    finally:
        db.close()

    # 提取 tone_distribution
    tone_dist = None
    if isinstance(final_report, dict):
        results = final_report.get('results', [])
        for r in results:
            td = (r.get('data') or {}).get('tone_distribution') or (r.get('preview') or {}).get('data')
            if isinstance(td, list) and td:
                tone_dist = td
                break
        if not tone_dist:
            cols = (final_report.get('statistics') or {}).get('colors', [])
            if cols:
                tone_dist = cols

    if not tone_dist:
        raise HTTPException(status_code=422, detail="该报告不包含色调分析数据，请确认已执行 analyze_tone")

    # 创建临时 AgentTools 实例，将 tone_analysis 写入缓存，再调用工具
    from backend.services.agent_tools import AgentTools
    tools = AgentTools()
    tools.data_cache['tone_analysis'] = tone_dist

    result = await tools.generate_color_schemes()
    if not result.get('success'):
        raise HTTPException(status_code=500, detail=result.get('error', '配色生成失败'))

    return result['data']  # 列表： [{"name": "...", "colors": ["#...", ...]}, ...]


# ============================================================
# AI 配色生成（直接上传报告文件，不需要 session）
# ============================================================

@router.post("/color-schemes/upload")
async def generate_color_schemes_from_file(
    file: UploadFile = File(...),
    authorization: Optional[str] = Header(None)
):
    """直接上传产品报告文件，提取文本内容后调用 LLM 生成 3 组配色方案"""
    _get_user_id(authorization)

    content_bytes = await file.read()
    filename = file.filename or ""
    text = ""

    try:
        if filename.lower().endswith(".pdf"):
            import io
            try:
                import pypdf
                reader = pypdf.PdfReader(io.BytesIO(content_bytes))
                text = "\n".join(page.extract_text() or "" for page in reader.pages)
            except ImportError:
                # fallback: 尝试 pdfminer
                from pdfminer.high_level import extract_text_to_fp
                from pdfminer.layout import LAParams
                out = io.StringIO()
                extract_text_to_fp(io.BytesIO(content_bytes), out, laparams=LAParams())
                text = out.getvalue()
        elif filename.lower().endswith((".xlsx", ".xls")):
            import io
            import openpyxl
            wb = openpyxl.load_workbook(io.BytesIO(content_bytes), data_only=True)
            lines = []
            for ws in wb.worksheets:
                for row in ws.iter_rows(values_only=True):
                    line = "  ".join(str(c) for c in row if c is not None)
                    if line.strip():
                        lines.append(line)
            text = "\n".join(lines)
        else:
            # .doc/.docx 或其他，尝试 python-docx
            import io
            try:
                import docx
                doc = docx.Document(io.BytesIO(content_bytes))
                text = "\n".join(p.text for p in doc.paragraphs if p.text.strip())
            except Exception:
                text = content_bytes.decode("utf-8", errors="ignore")
    except Exception as e:
        raise HTTPException(status_code=422, detail=f"文件解析失败：{e}")

    if not text.strip():
        raise HTTPException(status_code=422, detail="无法从文件中提取文本内容")

    # 截取前 3000 字符避免超出 token 限制
    excerpt = text[:3000]

    from backend.services.langchain_service import langchain_service
    import json as _json
    import re as _re

    prompt = (
        "下面是一份美妆产品趋势报告的内容节选：\n"
        f"{excerpt}\n\n"
        "请根据以上报告内容，设计 3 组美妆口红配色方案，每组 3 个颜色。"
        "要求：1. 颜色必须符合报告中提到的色调趋势；"
        "2. 每个颜色都要用具体的 6 位十六进制 HEX（如 #CC3300）；"
        "3. 每组方案取一个美妆风格名（中文，4-6个字）；"
        "4. 返回纯 JSON 数组，不要加三反引号代码块。"
        '格式：[{"name": "方案名", "colors": ["#RRGGBB", "#RRGGBB", "#RRGGBB"]}]'
    )

    resp = await langchain_service.llm.ainvoke(prompt)
    raw = resp.content if hasattr(resp, 'content') else str(resp)
    clean = _re.sub(r'^```[\w]*\n?|\n?```$', '', raw.strip(), flags=_re.MULTILINE).strip()

    try:
        schemes = _json.loads(clean)
    except Exception:
        raise HTTPException(status_code=500, detail=f"LLM 返回格式解析失败：{clean[:200]}")

    if not isinstance(schemes, list) or not schemes:
        raise HTTPException(status_code=500, detail="LLM 未返回有效的配色方案")

    return schemes
